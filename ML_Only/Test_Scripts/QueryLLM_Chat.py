from llama_index.llms.gemini import Gemini
from llama_index.core.llms import ChatMessage
from google.generativeai.types import HarmCategory, HarmBlockThreshold
import os
import openpyxl
from tqdm import tqdm

def build_query(sentence, subject, object, predicate):
    query = ("Sentence:"
             + sentence + "\n"
            "Subject: " + subject + "\n"
             "Object: " + object + "\n"
             "Predicate: " + predicate + "\n")
    return query

def gemini(llm, query):
    messages = [
        ChatMessage(role="user", content="You are now an expert in the task of Open Information Extraction. Identify all possible"
                                         "subject, predicate and objects for any given sentence, given a list of options for subject, predicate and object."
                                         "Avoid triples with the same meaning as much as possible"
                                         "Present your response in the form of triples: (subject, predicate, object)"),
        ChatMessage(role="assistant", content="Sure! I can help you with that. Please provide me with a sentence, and I will extract the"
                                              "subject, predicate and object given the constraits presented in the form of a triplet."),
        ChatMessage(
            role="user", content=query
        ),
    ]
    resp = llm.chat(messages)
    return resp


GOOGLE_API_KEY = "AIzaSyD1CATa-F8cdTRmRiLiXU3DfZIjv_MJz-c"
os.environ["GOOGLE_API_KEY"] = GOOGLE_API_KEY
safety_settings={
        HarmCategory.HARM_CATEGORY_HATE_SPEECH: HarmBlockThreshold.BLOCK_NONE,
        HarmCategory.HARM_CATEGORY_HARASSMENT: HarmBlockThreshold.BLOCK_NONE,
        HarmCategory.HARM_CATEGORY_SEXUALLY_EXPLICIT: HarmBlockThreshold.BLOCK_NONE,
        HarmCategory.HARM_CATEGORY_DANGEROUS_CONTENT: HarmBlockThreshold.BLOCK_NONE,
    }
llm = Gemini(model="models/gemini-pro", safety_settings=safety_settings)

wb_obj = openpyxl.load_workbook("Benchmarks/Results/excel/Roots_Only/CaRB_Tests.xlsx")
sheet_obj = wb_obj.active
row = 1
col = 1
row_results = 0
col_results = 0

for i in tqdm(range(0, sheet_obj.max_row, 5)): #8 para testes
    sentence = sheet_obj.cell(row=row + i, column=col+2).value
    sent_id = sheet_obj.cell(row=row + i, column=col+1).value
    for j in range(1, 4, 1):
    #for j in range(2, 7, 2): for testing
        current_predict = []
        for k in range(1, 20, 2):
            cell = sheet_obj.cell(row=row + i + j , column=col + k).value
            if cell == None:
                break
            else:
                current_predict.append(cell)
        if j == 1: #if j == 2: for testing
            subject = ""
            for sub in current_predict:
                subject += sub + " // "
        elif j == 2: #if j == 4 for testing
            predicate = ""
            for sub in current_predict:
                predicate += sub + " // "
        elif j == 3: #if j == 6 for testing
            object = ""
            for sub in current_predict:
                object += sub + " // "

    query = build_query(sentence, subject, predicate, object)
    #response = interview(temperature=0.5, project_id="stellar-zoo-405011", location="europe-west4", query=query)
    print(query)
    response = gemini(llm, query)
    print(response)



