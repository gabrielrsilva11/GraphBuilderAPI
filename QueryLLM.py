import vertexai
from vertexai.language_models import TextGenerationModel
import openpyxl
import xlsxwriter as xls
import regex as re
import os
from llama_index.llms.gemini import Gemini
from google.generativeai.types import HarmCategory, HarmBlockThreshold
from tqdm import tqdm
import string
import itertools

def interview(
    temperature: float,
    project_id: str,
    location: str,
    query: str,
) -> str:
    """Ideation example with a Large Language Model"""

    vertexai.init(project=project_id, location=location)
    # TODO developer - override these parameters as needed:
    parameters = {
        "temperature": temperature,  # Temperature controls the degree of randomness in token selection.
        "max_output_tokens": 256,  # Token limit determines the maximum amount of text output.
        "top_p": 0.8,  # Tokens are selected from most probable to least until the sum of their probabilities equals the top_p value.
        "top_k": 40,  # A top_k of 1 means the selected token is the most probable among all tokens.
    }

    model = TextGenerationModel.from_pretrained("text-bison@002")
    response = model.predict(
        query,
        **parameters,
    )
    #print(f"Response from Model: {response.text}")

    return response.text

def gemini(llm, query):
    resp = llm.complete(query)
    #print(type(resp))
    #print(resp.text)
    return resp.text

def query_builder_roots(sentence, subject, predicate, object):
    # In case there was no prediction
    if type(subject) is None:
        subject = ''
    if type(predicate) is None:
        predicate = ''
    if type(object) is None:
        object = ''

    query = ("""You are an Open Information Extraction model trained to find the most suitable triples (subject, predicate, object) in a sentence using the options separated by //. These options are to be used as root words of your solution and you are allowed to expand upon them. You can select one or more combinations to form triples.
        This is a very important task. Extract triples from the following sentence given the previous rules:
     Sentence: """+sentence+"""
     Options:
        Subject: """+subject+"""
        Predicate: """+predicate+"""
        Object: """+object+"""
     Extract the best triples without repeating Subject: and Object:.
     It is mandatory that your reply takes the following form:
     Subject:
     Predicate:
     Object:""")
    return query


def query_builder_example(sentence, subject, predicate, object):
    # In case there was no prediction
    if type(subject) == None:
        subject = ''
    if type(predicate) == None:
        predicate = ''
    if type(object) == None:
        object = ''
    query = ("""This is a very important task, you are an Open Information Extraction model trained to find the most suitable triples (subject, predicate, object) in a sentence given options separated by //. You can select one or more combinations to form triples. The triples are both explicit and implicit and may require using co-reference. The subject, predicate and object can not be punctuation. Try to form the most atomic triples possible.
As an example of a good extraction we have the following sentence:

Sentence: He stayed for less than a year before being appointed Chief Constable of Kent in July 1946 .

Subject: He
Predicate: stayed
Object: for less than a year

Subject: He
Predicate: was appointed
Object: Chief Constable of Kent in July 1946

Now extract triples from this sentence and options:
     Sentence: """+sentence+"""
     Options:
        Subject: """+subject+"""
        Predicate: """+predicate+"""
        Object: """+object+"""
     Extract the best triples without repeating Subject: and Object:.
     It is mandatory that your reply takes the following form:
     Subject:
     Predicate:
     Object:""")
    return query


def query_builder(sentence, subject, predicate, object):
    #In case there was no prediction
    if type(subject) == None:
        subject = ''
    if type(predicate) == None:
        predicate = ''
    if type(object) == None:
        object = ''

    # query = ("""You are an Open Information Extraction model trained to find the most suitable triples (subject, predicate, object) in a sentence given options separated by //.
    #      This is a very important task and you can select one or more combinations to form triples. The triples are both explicit and implicit and may require additional NLP processing (Such as entity linking or co-references).
    #      Sentence: """+sentence+"""
    #      Extract the best triples without repeating Subject: and Object:.
    #      It is mandatory that your reply takes the following form:
    #      Subject:
    #      Predicate:
    #      Object:""")
    # query = ("""You are an Open Information Extraction model trained to find the most suitable triples (subject, predicate, object) in a sentence given options separated by //.
    #      This is a very important task and you can select one or more combinations to form triples. The triples are both explicit and implicit and may require
    #      Sentence: """+sentence+"""
    #      Options:
    #         Subject: """+subject+"""
    #         Predicate: """+predicate+"""
    #         Object: """+object+"""
    #      Extract the best triples without repeating Subject: and Object:.
    #      It is mandatory that your reply takes the following form:
    #      Subject:
    #      Predicate:
    #      Object:""")
    query = ("""You are an Open Information Extraction model trained to extract relational triples (subject, predicate, object) from a sentence given options separated by //. 
You can select several combinations to form a single triple.
In cases where there is no option given you can decide the best option from the sentence. Do not leave subject, predicate or object empty.
 Sentence: """+sentence+"""
 Options:
    Subject: """+subject+"""
    Predicate: """+predicate+"""
    Object: """+object+"""

 It is mandatory that your reply takes the following form:
 
 Subject:
 Predicate:
 Object:
""")
    return query

def jodie_query_builder(sentence, subject, predicate, object):
    #In case there was no prediction
    if type(subject) == None:
        subject = ''
    if type(predicate) == None:
        predicate = ''
    if type(object) == None:
        object = ''

    query = ("""Jodie W. Jenkins is an expert Open Information Extraction annotator and she said "There are several triples in this sentence: """+sentence+""". 
        We have several options for subject, predicate and object presented:
         Options:
            Subject: """+subject+"""
            Predicate: """+predicate+""" 
            Object: """+object+""" 
        The triples are both explicit and implicit. We have to make sure the triples are as atomic as possible."
         Answer as if you are Jodie W. Jenkins.
         It is mandatory that your reply takes the following form:
         Subject:
         Predicate:
         Object:""")
    return query

# def fetch_from_reply(to_fetch, reply):
#     reply_split = reply.split(to_fetch)[-1].split("\n")[0].replace(": ", "")
#
#     if "*" in reply_split:
#         to_replace = "**" + to_fetch + "**:"
#         reply_split.replace(to_replace, "")
#         reply_split.replace("**", "")
#     return reply_split.strip()

def fetch_from_reply(to_fetch, reply):
    regex = to_fetch+":(.*)"
    found = re.findall(regex, reply)
    return found

def eliminate_empty_triples(triple_member):
    list_to_remove = []
    punct = string.punctuation
    for i in range(0, len(triple_member)):
        if triple_member[i] in punct or triple_member[i] == ' ' or triple_member[i] == '':
            list_to_remove.append(i)
    return list_to_remove


#Gemini Related
GOOGLE_API_KEY = "AIzaSyD1CATa-F8cdTRmRiLiXU3DfZIjv_MJz-c"
os.environ["GOOGLE_API_KEY"] = GOOGLE_API_KEY
safety_settings={
        HarmCategory.HARM_CATEGORY_HATE_SPEECH: HarmBlockThreshold.BLOCK_NONE,
        HarmCategory.HARM_CATEGORY_HARASSMENT: HarmBlockThreshold.BLOCK_NONE,
        HarmCategory.HARM_CATEGORY_SEXUALLY_EXPLICIT: HarmBlockThreshold.BLOCK_NONE,
        HarmCategory.HARM_CATEGORY_DANGEROUS_CONTENT: HarmBlockThreshold.BLOCK_NONE,
    }
llm = Gemini(model="models/gemini-pro", safety_settings=safety_settings)


triples_extracted = 0

#Ficheiro dos dados
#wb_obj = openpyxl.load_workbook("Results/Validation/Processed_GridSearch_PT_Teste.xlsx")
wb_obj = openpyxl.load_workbook("Benchmarks/Results/excel/Roots_Only/CaRB_Tests.xlsx")
sheet_obj = wb_obj.active
row = 1
col = 1
row_results = 0
col_results = 0

#Ficheiro a criar
#workbook = xls.Workbook('Results/Validation/Processed_LLM_TestQuery_PT_Formato.xlsx')
workbook = xls.Workbook('Benchmarks/Results/excel/Roots_Only/CaRB_Processed_Tests2_Gemini.xlsx')
worksheet = workbook.add_worksheet()
worksheet.write(row_results, col_results, "Sent_ID")
worksheet.write(row_results, col_results + 1, "Sentence")
worksheet.write(row_results, col_results + 2, "Subject")
worksheet.write(row_results, col_results + 3, "Predicate")
worksheet.write(row_results, col_results + 4, "Object")
worksheet.write(row_results, col_results + 5, "EVAL_Triple")
worksheet.write(row_results, col_results + 6, "EVAL_Subject")
worksheet.write(row_results, col_results + 7, "EVAL_Predicate")
worksheet.write(row_results, col_results + 8, "EVAL_Object")
row_results = 1
for i in tqdm(range(0, sheet_obj.max_row, 5)): #8 para testes
    sentence = sheet_obj.cell(row=row + i, column=col+2).value
    sent_id = sheet_obj.cell(row=row + i, column=col+1).value
    for j in range(1, 4, 1):
    #for j in range(2, 7, 2): for testing
        current_predict = []
        for k in range(1, 20, 2):
            cell = sheet_obj.cell(row=row + i + j , column=col + k).value
            if cell == None:
                break
            else:
                current_predict.append(cell)
        if j == 1: #if j == 2: for testing
            subject = ""
            for sub in current_predict:
                subject += sub + " // "
        elif j == 2: #if j == 4 for testing
            predicate = ""
            for sub in current_predict:
                predicate += sub + " // "
        elif j == 3: #if j == 6 for testing
            object = ""
            for sub in current_predict:
                object += sub + " // "

    query = query_builder(sentence, subject, predicate, object)
    #response = interview(temperature=0.5, project_id="stellar-zoo-405011", location="europe-west4", query=query)
    response = gemini(llm, query)
    # print(query)
    # print(response)

    subject_reply = fetch_from_reply("Subject", response)
    predicate_reply = fetch_from_reply("Predicate", response)
    object_reply = fetch_from_reply("Object", response)
    subject_eliminate = eliminate_empty_triples(subject_reply)
    predicate_eliminate = eliminate_empty_triples(predicate_reply)
    object_eliminate = eliminate_empty_triples(object_reply)
    joined_eliminate = set(itertools.chain(subject_eliminate, predicate_eliminate, object_eliminate))

    for idx in sorted(joined_eliminate, reverse=True):
        del subject_reply[idx]
        del predicate_reply[idx]
        del object_reply[idx]

    worksheet.write(row_results, col_results, sent_id)
    worksheet.write(row_results, col_results+1, sentence)
    row_results+=1
    current_row_results = row_results
    for k in range(0, len(subject_reply)):
        # if k == 0:
        worksheet.write(current_row_results, col_results + 2, subject_reply[k])
        current_row_results += 1
    current_row_results = row_results
    for k in range(0, len(predicate_reply)):
        worksheet.write(current_row_results, col_results + 3, predicate_reply[k])
        current_row_results += 1
    current_row_results = row_results
    for k in range(0, len(object_reply)):
        worksheet.write(current_row_results, col_results + 4, object_reply[k])
        current_row_results += 1
    rows_processed = max([len(subject_reply), len(predicate_reply), len(object_reply)])
    for k in range(row_results, row_results+rows_processed):
        worksheet.write(k, col_results, sent_id)

    row_results += rows_processed
        # else:
        #     worksheet.write(row_results + 1, col_results + 2, subject_reply[k])
    # worksheet.write(row_results+2, col_results, "Predicate")
    # for k in range(0, len(predicate_reply)):
    #     if k == 0:
    #         worksheet.write(row_results + 2, col_results + 2, predicate_reply[k])
    #     else:
    #         worksheet.write(row_results + 2, col_results + 1 + k*2, predicate_reply[k])
    # worksheet.write(row_results+3, col_results, "Object")
    # for k in range(0, len(object_reply)):
    #     if k == 0:
    #         worksheet.write(row_results + 3, col_results + 1 + k, object_reply[k])
    #     else:
    #         worksheet.write(row_results + 3, col_results + 1 + k*2, object_reply[k])
    triples_extracted += max([len(subject_reply), len(predicate_reply), len(object_reply)])
    #row_results += 5

worksheet.write(0, 18, "Triples Extracted:")
worksheet.write(0, 19, triples_extracted)
workbook.close()
